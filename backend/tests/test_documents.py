from __future__ import annotations

import io
from typing import Any

import httpx
import pytest
from docx import Document as DocxDocument

from fabrica.api.app import create_app
from fabrica.documents.extract import UnsupportedDocument, extract_text
from fabrica.pipeline.runner import InlineRunner
from tests.helpers import detail, who


def pdf_with(text: str) -> bytes:
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R"
        b" /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, 1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(
        b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF" % (len(objects) + 1, xref)
    )
    return out.getvalue()


def docx_with(paragraph: str, cells: list[str]) -> bytes:
    document = DocxDocument()
    document.add_paragraph(paragraph)
    table = document.add_table(rows=1, cols=len(cells))
    for cell, value in zip(table.rows[0].cells, cells, strict=True):
        cell.text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_extracts_pdf_docx_and_text() -> None:
    assert "Campos sociedad y fecha" in extract_text(
        "spec.pdf", pdf_with("Campos sociedad y fecha")
    )
    text = extract_text("cuestionario.docx", docx_with("Filtro por fecha", ["Campo", "Tipo"]))
    assert "Filtro por fecha" in text and "Campo | Tipo" in text
    assert extract_text("acta.txt", "Reunión: sí".encode("cp1252")) == "Reunión: sí"


def test_rejects_unsupported_or_empty_files() -> None:
    with pytest.raises(UnsupportedDocument):
        extract_text("foto.png", b"\x89PNG")
    with pytest.raises(UnsupportedDocument):
        extract_text("vacio.txt", b"   ")


@pytest.fixture
async def api() -> Any:
    app = create_app()
    runner = InlineRunner()
    app.state.runner = runner
    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
        yield client, runner


async def test_create_requirement_with_files(api: Any) -> None:
    client, runner = api
    resp = await client.post(
        "/api/requirements/with-files",
        data={
            "title": "Reporte de facturas",
            "description": "Listar facturas por fecha.",
            "kinds": ["especificacion", "cuestionario"],
        },  # fmt: skip
        files=[
            ("files", ("spec.pdf", pdf_with("Campos sociedad documento fecha"), "application/pdf")),
            (
                "files",
                (
                    "preguntas.docx",
                    docx_with("Filtro por fecha", ["A", "B"]),
                    "application/octet-stream",
                ),
            ),
        ],
        headers=who("ana", "funcional"),
    )
    assert resp.status_code == 201
    req_id = resp.json()["id"]
    await runner.wait_idle()
    docs = (
        await client.get(f"/api/requirements/{req_id}/documents", headers=who("ana", "funcional"))
    ).json()
    assert [(d["name"], d["kind"]) for d in docs] == [
        ("spec", "especificacion"),
        ("preguntas", "cuestionario"),
    ]
    assert (await detail(client, req_id))["requirement"]["stage"] == "aprobacion_cliente"


async def test_unsupported_upload_is_rejected(api: Any) -> None:
    client, _ = api
    resp = await client.post(
        "/api/requirements/with-files",
        data={"title": "Reporte", "description": "Descripción suficiente."},
        files=[("files", ("foto.png", b"\x89PNG", "image/png"))],
        headers=who("ana", "funcional"),
    )
    assert resp.status_code == 415
