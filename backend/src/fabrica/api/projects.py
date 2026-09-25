from __future__ import annotations

import io
from typing import Any

import openpyxl
from fastapi import APIRouter, HTTPException
from fastapi.responses import Response
from sqlalchemy import select

from fabrica.api.deps import Who
from fabrica.blackboard.service import Board
from fabrica.db.models import BillingSnapshot, Capability, Project, Requirement
from fabrica.db.session import session_scope
from fabrica.domain.schemas import (
    BillingSnapshotIn,
    BillingSnapshotOut,
    CapabilityIn,
    CapabilityOut,
    ProjectIn,
    ProjectOut,
)

router = APIRouter(prefix="/api/projects", tags=["proyectos"])

STAGE_WEIGHTS: dict[str, float] = {
    "recepcion": 0.10,
    "diseno": 0.30,
    "aprobacion_cliente": 0.40,
    "construccion": 0.60,
    "revision_abap": 0.80,
    "uat": 0.90,
    "cerrado": 1.00,
    "desestimado": 0.00,
}


@router.get("", response_model=list[ProjectOut])
async def list_projects(who: Who) -> list[Project]:
    async with session_scope() as s:
        rows = await s.scalars(select(Project).order_by(Project.name.asc()))
        return list(rows.all())


@router.post("", response_model=ProjectOut, status_code=201)
async def create_project(data: ProjectIn, who: Who) -> Project:
    async with session_scope() as s:
        project = Project(
            client_id=data.client_id,
            name=data.name,
            code=data.code.upper(),
            sap_system=data.sap_system,
        )
        s.add(project)
        await s.flush()
        return project


@router.get("/{project_id}/capabilities", response_model=list[CapabilityOut])
async def list_capabilities(project_id: int, who: Who) -> list[Capability]:
    async with session_scope() as s:
        rows = await s.scalars(
            select(Capability)
            .where(Capability.project_id == project_id)
            .order_by(Capability.name.asc())
        )
        return list(rows.all())


@router.post("/{project_id}/capabilities", response_model=CapabilityOut, status_code=201)
async def create_capability(project_id: int, data: CapabilityIn, who: Who) -> Capability:
    async with session_scope() as s:
        cap = Capability(project_id=project_id, name=data.name, code=data.code.upper())
        s.add(cap)
        await s.flush()
        return cap


@router.get("/{project_code}/snapshots", response_model=list[BillingSnapshotOut])
async def list_snapshots(project_code: str, who: Who) -> list[BillingSnapshot]:
    async with session_scope() as s:
        rows = await s.scalars(
            select(BillingSnapshot)
            .where(BillingSnapshot.project == project_code)
            .order_by(BillingSnapshot.id.desc())
        )
        return list(rows.all())


@router.post("/{project_code}/snapshots", response_model=BillingSnapshotOut, status_code=201)
async def create_snapshot(project_code: str, data: BillingSnapshotIn, who: Who) -> BillingSnapshot:
    async with session_scope() as s:
        board = Board(s)
        rows = await s.scalars(select(Requirement).where(Requirement.project == project_code))
        reqs = list(rows.all())

        items: list[dict[str, Any]] = []
        total_hours = 0.0
        billed_hours = 0.0

        for r in reqs:
            estimate = await board.latest_estimate(r.id)
            hours = estimate.hours_total if estimate else 0.0
            weight = STAGE_WEIGHTS.get(r.stage, 0.0)
            earned = hours * weight
            total_hours += hours
            billed_hours += earned
            items.append(
                {
                    "requirement_id": r.id,
                    "title": r.title,
                    "capability": r.capability or "-",
                    "stage": r.stage,
                    "state": r.state,
                    "hours": hours,
                    "weight": weight,
                    "billed": earned,
                }
            )

        snapshot = BillingSnapshot(
            project=project_code,
            created_by=who.user,
            hours_total=total_hours,
            hours_billed=billed_hours,
            items=items,
        )
        s.add(snapshot)
        await s.flush()
        return snapshot


@router.get("/{project_code}/export/billing")
async def export_billing_excel(project_code: str, who: Who) -> Response:
    async with session_scope() as s:
        board = Board(s)
        rows = await s.scalars(
            select(Requirement)
            .where(Requirement.project == project_code)
            .order_by(Requirement.id.asc())
        )
        reqs = list(rows.all())
        if not reqs:
            raise HTTPException(404, f"No hay requisitos para el proyecto {project_code}")

        snap_rows = await s.scalars(
            select(BillingSnapshot)
            .where(BillingSnapshot.project == project_code)
            .order_by(BillingSnapshot.id.desc())
            .limit(1)
        )
        last_snap = snap_rows.first()

        wb = openpyxl.Workbook()
        ws_summary = wb.active
        ws_summary.title = "Resumen"

        ws_summary.append(["Proyecto", project_code])
        ws_summary.append(["Generado por", who.user])
        ws_summary.append([])
        ws_summary.append(
            ["Capacidad", "Requisitos", "Horas Totales", "Horas Facturables", "% Avance"]
        )

        cap_totals: dict[str, dict[str, float]] = {}
        items_detail = []

        total_h = 0.0
        total_b = 0.0

        for r in reqs:
            cap = r.capability or "GENERAL"
            estimate = await board.latest_estimate(r.id)
            hours = estimate.hours_total if estimate else 0.0
            weight = STAGE_WEIGHTS.get(r.stage, 0.0)
            earned = hours * weight

            if cap not in cap_totals:
                cap_totals[cap] = {"count": 0, "hours": 0.0, "billed": 0.0}
            cap_totals[cap]["count"] += 1
            cap_totals[cap]["hours"] += hours
            cap_totals[cap]["billed"] += earned

            total_h += hours
            total_b += earned

            items_detail.append(
                [
                    r.id,
                    r.title,
                    cap,
                    r.stage,
                    r.state,
                    hours,
                    f"{int(weight * 100)}%",
                    earned,
                ]
            )

        for cap, val in sorted(cap_totals.items()):
            pct = f"{(val['billed'] / val['hours'] * 100):.1f}%" if val["hours"] > 0 else "0.0%"
            ws_summary.append([cap, int(val["count"]), val["hours"], val["billed"], pct])

        summary_pct = f"{(total_b / total_h * 100):.1f}%" if total_h > 0 else "0.0%"
        ws_summary.append(["TOTAL", len(reqs), total_h, total_b, summary_pct])

        ws_detail = wb.create_sheet(title="Detalle Requisitos")
        ws_detail.append(
            [
                "# ID",
                "Título",
                "Capacidad",
                "Etapa",
                "Estado",
                "Horas Plan",
                "% Avance Hito",
                "Horas Facturables",
            ]
        )
        for row in items_detail:
            ws_detail.append(row)

        ws_delta = wb.create_sheet(title="Delta Facturación")
        prev_billed = last_snap.hours_billed if last_snap else 0.0
        delta = total_b - prev_billed
        ws_delta.append(["Horas Facturables Acumuladas Actuales", total_b])
        ws_delta.append(["Horas Facturables en Última Foto Congelada", prev_billed])
        ws_delta.append(["Delta a Facturar este Período", delta])

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        filename = f"facturacion_{project_code.lower()}.xlsx"
        return Response(
            content=buf.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )
