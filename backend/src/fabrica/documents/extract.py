from __future__ import annotations

import io
from pathlib import PurePath

from docx import Document as DocxDocument
from pypdf import PdfReader
from pypdf.errors import PdfReadError

MAX_BYTES = 20 * 1024 * 1024
TEXT_SUFFIXES = {".txt", ".md", ".csv", ".json", ".xml", ".abap", ".sql", ".vtt", ".srt"}


class UnsupportedDocument(ValueError): ...


def _pdf(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
    except PdfReadError as exc:
        raise UnsupportedDocument(f"PDF ilegible: {exc}") from exc
    return "\n\n".join(p.strip() for p in pages if p.strip())


def _docx(data: bytes) -> str:
    document = DocxDocument(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _text(data: bytes) -> str:
    for encoding in ("utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise UnsupportedDocument("Texto con codificación desconocida")


def extract_text(filename: str, data: bytes) -> str:
    if len(data) > MAX_BYTES:
        raise UnsupportedDocument(f"{filename} supera {MAX_BYTES // (1024 * 1024)} MB")
    suffix = PurePath(filename).suffix.lower()
    if suffix == ".pdf":
        text = _pdf(data)
    elif suffix == ".docx":
        text = _docx(data)
    elif suffix in TEXT_SUFFIXES:
        text = _text(data)
    else:
        raise UnsupportedDocument(f"Formato no soportado: {suffix or filename}")
    if not text.strip():
        raise UnsupportedDocument(f"{filename} no contiene texto extraíble")
    return text.strip()


def document_name(filename: str) -> str:
    return PurePath(filename).stem[:180] or "documento"
